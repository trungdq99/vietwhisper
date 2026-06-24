import io
import zipfile
import docx

def format_timestamp(seconds: float) -> str:
    if seconds < 0:
        raise ValueError("Thời gian không được âm")
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = int(seconds % 60)
    ms = int((seconds % 1) * 1000)
    return f"[{h:02d}:{m:02d}:{s:02d}.{ms:03d}]"

def export_to_markdown(segments: list, include_timestamps: bool) -> io.BytesIO:
    bio = io.BytesIO()
    lines = []
    for seg in segments:
        text = seg.get("text", "").strip()
        if include_timestamps:
            ts = format_timestamp(seg.get("start", 0.0))
            lines.append(f"**{ts}** {text}")
        else:
            lines.append(text)
    content = "\n\n".join(lines)
    bio.write(content.encode("utf-8"))
    bio.seek(0)
    return bio

def export_to_docx(segments: list, include_timestamps: bool) -> io.BytesIO:
    doc = docx.Document()
    doc.add_heading("Kết quả chuyển đổi âm thanh", level=1)
    
    if include_timestamps:
        table = doc.add_table(rows=1, cols=2)
        try:
            table.style = 'Light Shading Accent 1'
        except KeyError:
            table.style = 'Table Grid' # Fallback style
            
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Mốc thời gian'
        hdr_cells[1].text = 'Văn bản'
        
        for seg in segments:
            row_cells = table.add_row().cells
            row_cells[0].text = format_timestamp(seg.get("start", 0.0))
            row_cells[1].text = seg.get("text", "").strip()
    else:
        for seg in segments:
            doc.add_paragraph(seg.get("text", "").strip())
            
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio

def export_to_zip(files: dict) -> io.BytesIO:
    zip_bio = io.BytesIO()
    with zipfile.ZipFile(zip_bio, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files.items():
            zf.writestr(name, data.getvalue())
    zip_bio.seek(0)
    return zip_bio
